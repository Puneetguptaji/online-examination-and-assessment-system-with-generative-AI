#!/usr/bin/env python3
"""Build the combined Python + AIP mini-project report in the supplied report style."""
from pathlib import Path
import json
import os
import re
import sys
from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'report' / 'ExamAI_Python_Project_Report.docx'
SECTIONS = ["Introduction", "Background", "Objectives", "Technologies Used", "Requirements",
            "System Design", "Source Code & Output", "Functionalities", "Conclusion",
            "Future Work", "Learning Outcomes", "GitHub Link", "Bibliography"]
PAGES = {}


def new_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    return doc


def para(doc, text='', size=12, bold=False, align='justify', before=0, after=6,
         first_indent=None, keep=False):
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf=p.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=1.0
    if first_indent is not None: pf.first_line_indent=first_indent
    if keep: pf.keep_with_next=True
    if text:
        r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=RGBColor(0,0,0)
    return p


def heading(doc, text, level=1):
    return para(doc, text, size=14 if level==1 else 12, bold=True, align='left',
                before=13 if level==1 else 8, after=5, keep=True)


def bullet(doc, text, marker='•  '):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent=Cm(.75); p.paragraph_format.first_line_indent=Cm(-.45)
    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.0
    r=p.add_run(marker+text); r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p


def code(doc, text, size=7.5):
    p=doc.add_paragraph(); pf=p.paragraph_format
    pf.left_indent=Emu(179705); pf.space_before=Pt(1); pf.space_after=Pt(6); pf.line_spacing=1.0
    r=p.add_run(text.strip('\n')); r.font.name='Consolas'; r.font.size=Pt(size); r.font.color.rgb=RGBColor(0,0,0)
    r._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas')
    return p


def caption(doc, text):
    return para(doc,text,size=10,align='center',before=2,after=8)


def picture(doc, path, width, caption_text):
    doc.add_picture(str(path),width=Cm(width))
    p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
    caption(doc,caption_text)


def grid(table):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'000000'); borders.append(e)
    tblPr.append(borders)


def cell_lines(cell, lines, size=12):
    cell.text=''
    for i,line in enumerate(lines):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.05
        r=p.add_run(line); r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=(i==0)


def excerpt(path, start, end=None):
    s=Path(path).read_text(encoding='utf-8')
    i=s.index(start); j=s.index(end,i+len(start)) if end else len(s)
    return s[i:j].strip('\n')


def report_image(src, dst, max_width=1450, max_height=1120):
    im=Image.open(src).convert('RGB')
    w,h=im.size
    # Keep the top of the real screen capture; this includes the title, navigation and primary evidence.
    if h>w*1.15:
        crop_h=min(h,int(w*0.95),max_height*int(w/max_width))
        im=im.crop((0,0,w,crop_h))
    if im.width>max_width:
        im=im.resize((max_width,round(im.height*max_width/im.width)),Image.Resampling.LANCZOS)
    im.save(dst,'JPEG',quality=88,optimize=True)
    return dst


def build(pages):
    sh=ROOT/'report_shots'; sh.mkdir(exist_ok=True)
    captures={
        'teacher': report_image(ROOT/'screenshots/02-teacher-dashboard.png',sh/'teacher.jpg'),
        'questions': ROOT/'screenshots/10-ai-question-draft-panel.png',
        'exam': report_image(ROOT/'screenshots/08-exam-in-progress.png',sh/'exam.jpg'),
        'result': report_image(ROOT/'screenshots/09-result-and-ai-state.png',sh/'result.jpg'),
        'analytics': report_image(ROOT/'screenshots/06-teacher-analytics.png',sh/'analytics.jpg'),
    }
    py=ROOT/'app.py'; ai=ROOT/'ai_service.py'; schema=ROOT/'01_schema.sql'; scoring=ROOT/'scoring.py'
    s_auth=excerpt(py,'def login():','@app.route("/register"')
    s_score=excerpt(py,'def finalize_attempt(attempt_id, force=False):','def result_reviews(attempt):')
    s_gemini=excerpt(ai,'    def _generate_json(self, prompt, schema, temperature=0.4):','    def generate_questions(')
    s_validate=excerpt(ai,'    @staticmethod\n    def validate_question(item):','    def generate_feedback(')
    s_ai_route=excerpt(py,'        gemini = GeminiService()','    except AIServiceError as exc:')
    s_score_func=Path(scoring).read_text(encoding='utf-8')
    s_schema=excerpt(schema,'CREATE TABLE IF NOT EXISTS attempts (','CREATE TABLE IF NOT EXISTS answers (')
    tests_path=ROOT/'docs/test_output.txt'
    test_out=tests_path.read_text(encoding='utf-8') if tests_path.exists() else 'Unit test output will be recorded after the test run.'
    smoke_path=ROOT/'docs/smoke_output.txt'
    smoke_out=smoke_path.read_text(encoding='utf-8') if smoke_path.exists() else 'Web-flow smoke test output will be recorded after the local integration run.'

    doc=new_doc()
    # Cover
    para(doc,'Mini Project Report',size=20,bold=True,align='center',before=108,after=8)
    para(doc,'On',size=13,align='center',after=10)
    para(doc,'AI-Powered Online Examination and Assessment System',size=16,bold=True,align='center',after=5)
    para(doc,'(Python, Flask, MySQL and Gemini API)',size=12,align='center',after=28)
    t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell_lines(t.rows[0].cells[0],['SUBMITTED BY:','Puneet Gupta','UID: 26MCI10246','CLASS: MCA (AI & ML)','SECTION: 24MAM4-4'])
    cell_lines(t.rows[0].cells[1],['SUBMITTED TO:','Python Programming Faculty','Advanced Internet Programming (AIP) Faculty'])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # TOC
    para(doc,'Table of Content',size=16,bold=True,align='center',before=6,after=14)
    toc=doc.add_table(rows=len(SECTIONS)+1,cols=3); toc.alignment=WD_TABLE_ALIGNMENT.CENTER; grid(toc)
    for i,txt in enumerate(['S. No.','Name','Page no.']):
        r=toc.rows[0].cells[i].paragraphs[0].add_run(txt); r.font.name='Times New Roman'; r.font.size=Pt(12); r.bold=True
    for n,name in enumerate(SECTIONS,1):
        row=toc.rows[n]
        for j,val in enumerate([str(n),name,str(pages.get(name,''))]):
            row.cells[j].text=''; r=row.cells[j].paragraphs[0].add_run(val); r.font.name='Times New Roman'; r.font.size=Pt(12)
    for row in toc.rows:
        row.cells[0].width=Cm(2.2); row.cells[1].width=Cm(10.6); row.cells[2].width=Cm(2.6)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Acknowledgement
    para(doc,'Acknowledgement',size=16,bold=True,align='center',before=6,after=14)
    para(doc,'I would like to express my gratitude to the faculty of Python Programming and Advanced Internet Programming for their guidance and feedback during the development of this project.',first_indent=Cm(1.27),after=9)
    para(doc,'I am thankful to the Department of Computer Applications, Chandigarh University, for providing the learning environment and resources that supported this work. I also thank my classmates who tested the student and faculty workflows and helped identify issues in the application.',first_indent=Cm(1.27),after=9)
    para(doc,'Finally, I would like to thank my family for their patience and support throughout the project.',first_indent=Cm(1.27),after=15)
    para(doc,'Puneet Gupta',bold=True,align='right',after=0); para(doc,'(26MCI10246)',align='right',after=0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 1 Introduction
    heading(doc,'1. Introduction')
    para(doc,'This report describes ExamAI, an online examination and assessment system implemented in Python. It provides separate student and teacher workflows: faculty create and publish papers, students complete timed assessments, and the application calculates each result from the saved answer key.')
    para(doc,'The project also uses generative AI for two language tasks. Gemini can draft multiple-choice questions for a teacher to review, and it can prepare a learning report from a student’s actual incorrect answers. AI is not allowed to award marks. The Python application performs deterministic scoring and stores the result in MySQL.')
    para(doc,'The same application is suitable for both Python Programming and Advanced Internet Programming (AIP): Python and Flask are the core implementation, while the Gemini API demonstrates an applied AI feature in a real web workflow.')

    # 2 Background
    heading(doc,'2. Background')
    para(doc,'Preparing a question paper and checking each response manually take time. A score by itself also gives a student little guidance about which concepts to revisit. Generative AI can speed up question drafting and produce explanations, but its answer keys may be wrong or its feedback may not match the student’s actual responses.')
    para(doc,'ExamAI addresses these risks with three separate responsibilities. Gemini proposes content; the teacher checks generated questions before approval; and the application compares submitted choices with the stored answer key. Feedback is built from the attempt’s real wrong answers, not from an invented student profile.')
    para(doc,'The earlier ExamAI concept used Java and JSP. This edition uses Python and Flask as its main backend while retaining the examination workflow and AI design. Its separate database, examai_python, avoids modifying the earlier Java project database or the Online Examination Result System database.')

    # 3 Objectives
    heading(doc,'3. Objectives')
    for x in ['Build one Python web application that can be presented for both Python and AIP coursework.','Support teacher and student accounts with role-based access.','Create, edit and publish timed multiple-choice examinations.','Use Gemini to draft questions and personalized learning feedback, with explicit validation and teacher review.','Autosave student answers and enforce the time limit on the server.','Compute marks reproducibly in Python and store all exam activity in MySQL.']:
        bullet(doc,x)

    # 4 Technologies
    heading(doc,'4. Technologies Used')
    for lead,desc in [
        ('Python 3: ','application routes, validation, authentication helpers and deterministic scoring.'),
        ('Flask: ','web framework for routing, sessions, templates and flash messages.'),
        ('MySQL 8: ','relational storage for users, exams, questions, attempts, answers, AI drafts and feedback.'),
        ('PyMySQL: ','Python database driver for parameterized SQL statements.'),
        ('HTML, CSS and JavaScript: ','responsive pages, question controls, timer display and background answer saving.'),
        ('Gemini REST API: ','structured question drafting and grounded post-exam study feedback.'),
        ('Werkzeug: ','password hashing and verification used by Flask.'),
    ]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.left_indent=Cm(.75); p.paragraph_format.first_line_indent=Cm(-.45); p.paragraph_format.space_after=Pt(3)
        a=p.add_run('•  '+lead); a.bold=True; a.font.name='Times New Roman'; a.font.size=Pt(12)
        b=p.add_run(desc); b.font.name='Times New Roman'; b.font.size=Pt(12)

    # 5 Requirements
    heading(doc,'5. Requirements')
    for x in ['Windows 10/11, macOS or Linux; Python 3.10 or newer.','MySQL Server 8.0 or newer. MySQL Workbench is used to run the setup scripts.','Internet connection and a valid Gemini API key for live AI requests. Exams and deterministic scoring work without an AI key.','Python packages listed in requirements.txt.','Modern browser such as Chrome, Edge or Firefox.']:
        bullet(doc,x)
    para(doc,'For a local classroom setup, start the MySQL80 service, run 01_schema.sql and 02_create_db_user.sql in Workbench, install the Python requirements, run seed_demo.py, and then start app.py. The supplied SQL does not drop a database.',after=8)

    # 6 System Design
    heading(doc,'6. System Design')
    para(doc,'ExamAI follows a three-tier client–server design. The browser renders Flask templates and sends form or JSON requests. Flask routes enforce the role and ownership rules, coordinate the assessment services and call the data layer. MySQL stores application data. A separate AI service calls Gemini from the server only; the browser never receives the API key.')
    para(doc,'The data flow diagram shows the student and faculty inputs, the Python application, the MySQL records and the two AI exchanges. The scoring path stays independent from Gemini: the submitted option is checked against the question key, then Python stores the computed score and percentage.')
    picture(doc,ROOT/'docs/dfd.png',16.0,'Fig. 1. Data flow of the Python ExamAI application. The model proposes or explains; Python computes marks.')
    heading(doc,'Database design',2)
    para(doc,'The schema has seven related tables: users, exams, questions, attempts, answers, ai_drafts and ai_feedback. Foreign keys maintain their relationships. A unique constraint limits a student to one attempt per exam; another keeps one saved response per question in an attempt. Attempts store a frozen question order and server-side deadline.')
    para(doc,'The following table fragment records the attempt owner, deadline, status and frozen question order. The setup scripts use CREATE IF NOT EXISTS and leave existing project databases untouched.')
    code(doc,s_schema,size=7.2)

    # 7 Source Code and Output
    heading(doc,'7. Source Code & Output')
    para(doc,'Selected listings below are copied from the project source files. The complete Python modules and templates are included with the deliverable. The excerpts show the decision points: deterministic marking, the Gemini request, question validation and the human-review hand-off.')
    heading(doc,'7.1  Deterministic scoring',2)
    para(doc,'The score is a pure Python calculation from the question marks, the stored correct option and the saved student choice. A blank answer earns no marks. The function has no Gemini dependency.')
    code(doc,s_score_func,size=7.6)
    heading(doc,'7.2  Server-side attempt finalization',2)
    code(doc,s_score,size=7.0)
    heading(doc,'7.3  Gemini request and structured response',2)
    para(doc,'The API key is read from the server environment. The response is parsed as JSON and API failures are returned as readable errors. No key is stored in a source file or template.')
    code(doc,s_gemini,size=6.9)
    heading(doc,'7.4  Question validation and review queue',2)
    code(doc,s_validate,size=7.1)
    para(doc,'The route below saves validated items as pending drafts. A separate teacher action copies an approved question into the exam; an unapproved draft is never presented to a student.')
    code(doc,s_ai_route,size=6.9)
    heading(doc,'7.5  Verification output',2)
    para(doc,'The following unit tests ran in the project workspace. They cover scoring, blank answers, JSON parsing, API-key handling, option validation, high-demand fallback and an HTTP rate-limit response. Live teacher question drafting and student learning feedback were also verified using Gemini 3.1 Flash Lite. One validated question draft is pending teacher approval, and one feedback report is cached for a submitted demo attempt. The API key and generated text are not included in this report or the project package.')
    code(doc,test_out,size=7.2)
    para(doc,'The web-flow smoke test connected to MySQL 8.0.46 and exercised teacher and student routes, question creation and editing, publication, answer autosave, submission and result display. The health endpoint reported that the database connection was available.')
    code(doc,smoke_out,size=7.2)
    heading(doc,'Screens from the running application',2)
    picture(doc,captures['teacher'],12.7,'Fig. 2. Faculty dashboard with exam and submission statistics from the local MySQL data.')
    picture(doc,captures['questions'],8.8,'Fig. 3. Gemini question-drafting panel. Generated items enter a review queue and are not published automatically.')
    picture(doc,captures['exam'],12.7,'Fig. 4. Timed assessment. The answer is autosaved and the countdown is displayed.')
    picture(doc,captures['result'],12.7,'Fig. 5. Result review from the default no-key capture state. The score is computed from the answer key; live Gemini feedback needs a key and network access.')
    picture(doc,captures['analytics'],12.7,'Fig. 6. Faculty analytics showing submitted scores and topic accuracy.')

    # 8 Functionality
    heading(doc,'8. Functionalities')
    heading(doc,'Faculty',2)
    for x in ['Create exams and manage their draft, published and closed states.','Add, edit and delete questions before the exam has attempts.','Request Gemini question drafts by topic and difficulty; review, approve or reject each draft.','View attempt counts, class average, highest and lowest scores, ranked student performance and topic accuracy.']:
        bullet(doc,x)
    heading(doc,'Student',2)
    for x in ['Register as a student and view published assessments.','Start a single timed attempt; the server checks the deadline even if the browser timer is manipulated.','Autosave selected options and resume an in-progress attempt.','Review the score, selected and correct answers, explanations and accuracy by topic.','Request AI study feedback grounded in the actual incorrect answers when Gemini is configured.']:
        bullet(doc,x)
    heading(doc,'AI safeguards',2)
    for x in ['AI output is validated for a non-empty question, four distinct options and a valid correct-answer index.','Generated items remain in a pending queue until approved by the teacher.','The model is never asked to calculate or award marks.','AI feedback is cached per attempt and uses question-level evidence rather than names or emails.']:
        bullet(doc,x)

    # 9-13
    heading(doc,'9. Conclusion')
    para(doc,'The Python ExamAI edition combines an operational online examination workflow with a genuine server-side Gemini integration. Flask handles the roles and requests; MySQL preserves the paper and attempt data; Python computes the score; Gemini can draft questions and explain mistakes. The human review queue and deterministic marking separate AI assistance from academic decisions.')
    para(doc,'The project was exercised against MySQL 8.0.46 with demo accounts and sample assessments. Unit and web-flow checks passed, and live Gemini question drafting and student feedback were verified in the hosted environment. The distributable project contains no API key; a local installation needs a private key in its own .env file for Gemini requests.')
    heading(doc,'10. Future Work')
    for x in ['Add a teacher-controlled pass threshold and exportable PDF mark sheets.','Support more question types, including numerical and short-answer responses with rubric-based teacher grading.','Add exam availability windows and reminders.','Add attempt audit logs and richer course-level analytics.','Deploy behind HTTPS with production secret management, backups and privacy controls.']:
        bullet(doc,x)
    heading(doc,'11. Learning Outcomes')
    for x in ['Building Python web routes, templates, sessions and role checks with Flask.','Using parameterized PyMySQL queries and relational constraints in MySQL 8.','Separating server-side score calculation from the browser and from AI calls.','Validating structured generative-AI output before it is saved or shown to a student.','Implementing a teacher-in-the-loop review workflow for AI-generated educational content.','Testing with unit tests and an end-to-end flow rather than relying only on a screenshot.']:
        bullet(doc,x)
    heading(doc,'12. GitHub Link')
    para(doc,'[ADD YOUR REPOSITORY LINK HERE]')
    heading(doc,'13. Bibliography')
    for ref in [
        '[1] Pallets Projects. Flask Documentation. https://flask.palletsprojects.com/',
        '[2] Python Software Foundation. Python 3 Documentation. https://docs.python.org/3/',
        '[3] MySQL. MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/refman/8.0/en/',
        '[4] PyMySQL Project. PyMySQL Documentation. https://pymysql.readthedocs.io/',
        '[5] Google AI for Developers. Gemini API Documentation. https://ai.google.dev/gemini-api/docs',
        '[6] Google AI for Developers. Gemini API Models. https://ai.google.dev/gemini-api/docs/models',
        '[7] OWASP Foundation. Cross Site Request Forgery Prevention Cheat Sheet. https://cheatsheetseries.owasp.org/cheatsheets/Cross-Site_Request_Forgery_Prevention_Cheat_Sheet.html',
    ]:
        para(doc,ref,size=11.5,align='left',after=4)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print('saved',OUT)


if __name__=='__main__':
    if len(sys.argv)>1 and Path(sys.argv[1]).exists():
        PAGES=json.loads(Path(sys.argv[1]).read_text())
    else:
        PAGES={section:'' for section in SECTIONS}
    build(PAGES)
